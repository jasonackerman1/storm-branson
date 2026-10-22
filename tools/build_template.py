"""Builds the one-tab spreadsheet ("Tracker") the dashboard reads, in two versions:
  template/Storm-Branson-Tracker-with-sample-data.xlsx  the roster plus SAMPLE money
      and sponsor rows (from js/sample-data.js) to test the connection. Every sample
      row's Notes starts with "SAMPLE"; the dashboard warns until they're deleted.
  template/Storm-Branson-Tracker-blank.xlsx  just the roster.

Run:  python3 tools/build_template.py

How the dashboard reads the tab:
  - a row with a #      = a player on the roster
  - a row with Amount   = money received (Player blank = General Fund)
  - a Sponsor row with a Status and no Amount = a company that's been asked
The # column is plain text so "?" survives: Google's CSV export blanks any value
whose type doesn't match the rest of its column.
"""
import json
import os
import subprocess
from datetime import datetime
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.join(HERE, '..', 'template')
SAMPLE_JS = os.path.join(HERE, '..', 'js', 'sample-data.js')
MAX_ROWS = 1000

COLUMNS = [
    # header, width, hover note
    ('Player', 22, 'The player this row is for. Leave blank for a donation to the whole team (General Fund).'),
    ('#', 6, 'Jersey number. Only on the roster rows at the top: a row with a # adds that player to the dashboard.'),
    ('Date', 12, 'The day the money came in.'),
    ('From', 28, 'Who it came from: a company, Grandma & Grandpa, a fundraiser table...'),
    ('Amount', 11, 'Dollars received. Leave blank on a sponsor row until their money actually arrives.'),
    ('Method', 11, 'Check, Venmo, Cash, Zelle, PayPal, or Other.'),
    ('Type', 12, 'Deposit = the $50 commitment deposit. Sponsor = a company. Family, Fundraiser (e.g. a table split between boys: one row per boy with his share), or Other.'),
    ('Status', 11, 'Sponsors only: Asked, Pledged, Paid, or Declined.'),
    ('Notes', 40, 'Anything else, e.g. "Pledged $150" or "Split with Jake ($125 total)".'),
]
HEAD_FILL = PatternFill('solid', fgColor='000000')
HEAD_FONT = Font(bold=True, color='F8B000')
ROSTER_FILL = PatternFill('solid', fgColor='FFF4D6')


def load_sample():
    js = ('global.window={};require(process.argv[1]);'
          'process.stdout.write(JSON.stringify(window.BRANSON_SAMPLE))')
    return json.loads(subprocess.check_output(['node', '-e', js, os.path.abspath(SAMPLE_JS)]))['rows']


def build(out, sample_rows, with_money):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tracker'
    for i, (name, width, note) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.comment = Comment(note, 'Storm')
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = 'A2'

    for r in range(2, MAX_ROWS + 1):
        ws.cell(row=r, column=2).number_format = '@'
        ws.cell(row=r, column=3).number_format = 'm/d/yyyy'
        ws.cell(row=r, column=5).number_format = '"$"#,##0.00'
        for col in (1, 4, 9):
            ws.cell(row=r, column=col).number_format = '@'

    rows = [r for r in sample_rows if r.get('number')]          # roster first
    if with_money:
        rows += [r for r in sample_rows if not r.get('number')]
    for i, r in enumerate(rows, start=2):
        roster = bool(r.get('number'))
        ws.cell(row=i, column=1, value=r.get('player') or None)
        ws.cell(row=i, column=2, value=r.get('number') or None)
        if r.get('date'):
            ws.cell(row=i, column=3, value=datetime.strptime(r['date'], '%m/%d/%Y'))
        ws.cell(row=i, column=4, value=r.get('from') or None)
        if r.get('amount'):
            ws.cell(row=i, column=5, value=float(r['amount']))
        ws.cell(row=i, column=6, value=r.get('method') or None)
        ws.cell(row=i, column=7, value=r.get('type') or None)
        ws.cell(row=i, column=8, value=r.get('status') or None)
        if roster:
            ws.cell(row=i, column=1).fill = ws.cell(row=i, column=2).fill = ROSTER_FILL
        else:
            note = r.get('notes') or ''
            ws.cell(row=i, column=9, value='SAMPLE - delete this row' + (' (' + note + ')' if note else ''))

    for rng, items in (('F2:F%d', 'Check,Venmo,Cash,Zelle,PayPal,Other'),
                       ('G2:G%d', 'Deposit,Sponsor,Family,Fundraiser,Other'),
                       ('H2:H%d', 'Asked,Pledged,Paid,Declined')):
        dv = DataValidation(type='list', formula1='"' + items + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng % MAX_ROWS)

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print('wrote', os.path.normpath(out), '-', len(rows), 'rows')


sample = load_sample()
build(os.path.join(TEMPLATE_DIR, 'Storm-Branson-Tracker-with-sample-data.xlsx'), sample, True)
build(os.path.join(TEMPLATE_DIR, 'Storm-Branson-Tracker-blank.xlsx'), sample, False)
